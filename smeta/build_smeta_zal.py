#!/usr/bin/env python3
"""Локальная смета на работы. Объект: Симагино. Зал: грунт, обои, покраска.

Итог зафиксирован: 50 000 ₽. Только работы, без материалов.
"""

from __future__ import annotations

from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.page import PageMargins


NAVY = "1B365D"
NAVY_DEEP = "12243F"
STEEL = "3D5A80"
GOLD = "C4A35A"
GOLD_SOFT = "E6D3A3"
CREAM = "FBF8F2"
SAND = "F3EBDD"
ROW_ALT = "F7F3EA"
WHITE = "FFFFFF"
INK = "1F2933"
MUTED = "5C6670"
LINE = "D7CDBD"
LINE_SOFT = "E8E0D4"
SUBTOTAL_BG = "E7EEF6"
GREEN = "2F6B4F"
TERRACOTTA = "A65D3F"
PALE_BLUE = "EEF3F8"

TARGET = 50_000
DOC_NO = "ЛС-СИМ-02"
OBJECT_NAME = "Симагино"
ROOM = "Зал"
DOC_DATE = date(2026, 10, 5)
OUT = Path(__file__).resolve().parent / "Smeta_raboty_Simagino_zal.xlsx"


def round_rub(n: float) -> int:
    return int(Decimal(str(n)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def money(n: float) -> str:
    return f"{int(round_rub(n)):,.0f}".replace(",", " ")


def rub_words(n: int) -> str:
    units = (
        ("", "", ""),
        ("тысяча", "тысячи", "тысяч"),
        ("миллион", "миллиона", "миллионов"),
    )
    ones_m = ["", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
    ones_f = ["", "одна", "две", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
    teens = [
        "десять", "одиннадцать", "двенадцать", "тринадцать", "четырнадцать",
        "пятнадцать", "шестнадцать", "семнадцать", "восемнадцать", "девятнадцать",
    ]
    tens = ["", "", "двадцать", "тридцать", "сорок", "пятьдесят", "шестьдесят", "семьдесят", "восемьдесят", "девяносто"]
    hundreds = ["", "сто", "двести", "триста", "четыреста", "пятьсот", "шестьсот", "семьсот", "восемьсот", "девятьсот"]

    def tri(num: int, feminine: bool) -> str:
        h, rest = divmod(num, 100)
        t, o = divmod(rest, 10)
        ones = ones_f if feminine else ones_m
        parts = []
        if h:
            parts.append(hundreds[h])
        if t == 1:
            parts.append(teens[o])
        else:
            if t:
                parts.append(tens[t])
            if o:
                parts.append(ones[o])
        return " ".join(parts)

    def form(num: int, forms: tuple[str, str, str]) -> str:
        n100 = num % 100
        n10 = num % 10
        if 11 <= n100 <= 19:
            return forms[2]
        if n10 == 1:
            return forms[0]
        if 2 <= n10 <= 4:
            return forms[1]
        return forms[2]

    n = int(n)
    if n == 0:
        return "Ноль рублей 00 копеек"
    chunks = []
    tmp = n
    while tmp:
        chunks.append(tmp % 1000)
        tmp //= 1000
    words = []
    for i, chunk in enumerate(chunks):
        if chunk == 0:
            continue
        w = tri(chunk, feminine=(i == 1))
        if i:
            w = f"{w} {form(chunk, units[i])}".strip()
        words.append(w)
    words.reverse()
    body = " ".join(words).strip()
    body = body[0].upper() + body[1:]
    return f"{body} {form(n, ('рубль', 'рубля', 'рублей'))} 00 копеек"


# Только три работы из задания. Объёмы договорные под итог 50 000 ₽.
SECTIONS: list[dict] = [
    {
        "num": "1",
        "title": "Отделка зала",
        "intro": "Помещение — зал. Материалы (грунт, обои, клей, краска) закупает заказчик.",
        "items": [
            {
                "code": "1.1",
                "name": "Грунтовка стен зала",
                "unit": "м²",
                "qty": 50,
                "price": 100,
                "note": "Один слой по стенам. Потолок не грунтуется",
            },
            {
                "code": "1.2",
                "name": "Поклейка обоев",
                "unit": "м²",
                "qty": 50,
                "price": 500,
                "note": "Метровые или стандартные. Подгонка рисунка — по факту обоев",
            },
            {
                "code": "1.3",
                "name": "Покраска",
                "unit": "м²",
                "qty": 40,
                "price": 500,
                "note": "Потолок и поверхности без обоев. Два слоя",
            },
        ],
    },
]

VOLUMES = [
    ("Стены зала под грунтовку", "50 м²", "п. 1.1", "За вычетом проёмов"),
    ("Стены под обои", "50 м²", "п. 1.2", "Та же площадь стен, что и под грунт"),
    ("Покраска", "40 м²", "п. 1.3", "Потолок и поверхности без обоев"),
]

MATERIALS = [
    ("Грунт", "Грунтовка под обои и под краску"),
    ("Обои", "Обои на площадь стен зала + запас 10–15%"),
    ("Клей", "Клей под выбранный тип обоев"),
    ("Краска", "Краска для потолка и поверхностей без обоев"),
    ("Расходники", "Валики, кисти, малярная лента, плёнка"),
]


def section_sum(section: dict) -> int:
    return sum(round_rub(item["qty"] * item["price"]) for item in section["items"])


def grand_total() -> int:
    return sum(section_sum(section) for section in SECTIONS)


def thin_border(color: str = LINE) -> Border:
    side = Side(style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def font(name: str = "Calibri", size: int = 11, bold: bool = False, color: str = INK, italic: bool = False) -> Font:
    return Font(name=name, size=size, bold=bold, color=color, italic=italic)


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def align(h: str = "left", v: str = "center", wrap: bool = True) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def apply_print(ws, landscape: bool, title: str) -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.45, right=0.45, top=0.65, bottom=0.55, header=0.25, footer=0.25)
    ws.oddHeader.left.text = f"&K1B365D{DOC_NO}  ·  {OBJECT_NAME}  ·  {ROOM}"
    ws.oddHeader.right.text = "&K5C6670только работы"
    ws.oddFooter.left.text = f"&K5C6670{title}"
    ws.oddFooter.center.text = "&K5C6670стр. &P из &N"
    ws.oddFooter.right.text = "&K5C667050 000 ₽"
    ws.print_options.horizontalCentered = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.view = "normal"


def set_widths(ws, widths: dict[str, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def paint(ws, row: int, c1: int, c2: int, fl=None, ft=None, al=None, bd=None, height: float | None = None) -> None:
    if height is not None:
        ws.row_dimensions[row].height = height
    for col in range(c1, c2 + 1):
        cell = ws.cell(row, col)
        if fl:
            cell.fill = fl
        if ft:
            cell.font = ft
        if al:
            cell.alignment = al
        if bd:
            cell.border = bd


def cream_sheet(ws, rows: int, cols: int) -> None:
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            ws.cell(r, c).fill = fill(CREAM)


def build_estimate(wb: Workbook) -> dict:
    ws = wb.active
    ws.title = "Смета"
    apply_print(ws, landscape=True, title="Локальная смета")
    set_widths(ws, {"A": 8, "B": 48, "C": 12, "D": 12, "E": 14, "F": 16, "G": 52})
    last_col = 7
    cream_sheet(ws, 28, last_col)

    paint(ws, 1, 1, last_col, fill(NAVY_DEEP), height=22)
    ws.merge_cells("A1:G1")
    ws["A1"] = f"ЛОКАЛЬНАЯ СМЕТА НА РАБОТЫ   ·   {DOC_NO}   ·   {OBJECT_NAME.upper()}   ·   {ROOM.upper()}"
    ws["A1"].font = font(size=12, bold=True, color=GOLD)
    ws["A1"].alignment = align("left", "center")
    ws["A1"].fill = fill(NAVY_DEEP)

    paint(ws, 2, 1, last_col, fill(GOLD), height=5)

    ws.merge_cells("A3:G3")
    ws["A3"] = "Только работы. Грунт, обои, клей и краска в итог не входят."
    ws["A3"].font = font(size=10, italic=True, color=MUTED)
    ws["A3"].alignment = align("left", "center")
    ws["A3"].fill = fill(CREAM)
    ws.row_dimensions[3].height = 18

    headers = ["№", "Наименование работ", "Ед. изм.", "Кол-во", "Расценка, ₽", "Сумма, ₽", "Что входит в строку"]
    header_row = 5
    ws.row_dimensions[header_row].height = 24
    for col, text in enumerate(headers, 1):
        cell = ws.cell(header_row, col, text)
        cell.font = font(size=10, bold=True, color=WHITE)
        cell.fill = fill(NAVY)
        cell.alignment = align("center", "center")
        cell.border = thin_border(NAVY)
    ws.freeze_panes = "A6"
    ws.print_title_rows = "1:5"

    row = 6
    section_rows = []
    money_format = "#,##0"
    for section in SECTIONS:
        paint(ws, row, 1, last_col, fill(STEEL), height=22)
        ws.cell(row, 1, section["num"]).font = font(size=12, bold=True, color=GOLD)
        ws.cell(row, 1).alignment = align("center")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col)
        ws.cell(row, 2, section["title"]).font = font(size=12, bold=True, color=WHITE)
        ws.cell(row, 2).alignment = align("left", "center")
        for col in range(1, last_col + 1):
            ws.cell(row, col).fill = fill(STEEL)
            ws.cell(row, col).border = thin_border(STEEL)
        row += 1

        paint(ws, row, 1, last_col, fill(SAND), height=20)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
        ws.cell(row, 1, section["intro"]).font = font(size=9, italic=True, color=STEEL)
        ws.cell(row, 1).alignment = align("left", "center")
        ws.cell(row, 1).fill = fill(SAND)
        row += 1

        first_item = row
        for item_i, item in enumerate(section["items"]):
            bg = WHITE if item_i % 2 == 0 else ROW_ALT
            ws.row_dimensions[row].height = 30
            values = [
                item["code"],
                item["name"],
                item["unit"],
                item["qty"],
                item["price"],
                f"=ROUND(D{row}*E{row},0)",
                item["note"],
            ]
            aligns = ["center", "left", "center", "center", "right", "right", "left"]
            for col, (value, how) in enumerate(zip(values, aligns), 1):
                cell = ws.cell(row, col, value)
                cell.font = font(size=11, bold=(col == 6), color=INK)
                cell.fill = fill(bg)
                cell.alignment = align(how, "center")
                cell.border = thin_border(LINE_SOFT)
            ws.cell(row, 2).font = font(size=11, bold=True, color=INK)
            ws.cell(row, 4).number_format = "0"
            ws.cell(row, 5).number_format = money_format
            ws.cell(row, 6).number_format = money_format
            row += 1
        last_item = row - 1

        paint(ws, row, 1, last_col, fill(SUBTOTAL_BG), height=22)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        ws.cell(row, 1, f"Итого по разделу {section['num']}.  {section['title']}")
        for col in range(1, 6):
            ws.cell(row, col).fill = fill(SUBTOTAL_BG)
            ws.cell(row, col).border = thin_border(LINE)
            ws.cell(row, col).font = font(size=10, bold=True, color=NAVY)
        ws.cell(row, 1).alignment = align("right", "center")
        sum_cell = ws.cell(row, 6, f"=SUM(F{first_item}:F{last_item})")
        sum_cell.font = font(size=12, bold=True, color=NAVY)
        sum_cell.fill = fill(SUBTOTAL_BG)
        sum_cell.alignment = align("right", "center")
        sum_cell.number_format = money_format
        sum_cell.border = thin_border(LINE)
        ws.cell(row, 7).fill = fill(SUBTOTAL_BG)
        ws.cell(row, 7).border = thin_border(LINE)
        section_rows.append(
            {
                "num": section["num"],
                "title": section["title"],
                "sum": section_sum(section),
                "cell": f"F{row}",
                "row": row,
            }
        )
        row += 1

    row += 1
    total_row = row
    paint(ws, total_row, 1, last_col, fill(NAVY), height=30)
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=5)
    ws.cell(total_row, 1, "ИТОГО РАБОТЫ")
    ws.cell(total_row, 1).font = font(size=14, bold=True, color=WHITE)
    ws.cell(total_row, 1).alignment = align("right", "center")
    parts = "+".join(section["cell"] for section in section_rows)
    total_cell = ws.cell(total_row, 6, f"={parts}")
    total_cell.font = font(size=16, bold=True, color=GOLD)
    total_cell.alignment = align("right", "center")
    total_cell.number_format = '#,##0" ₽"'
    for col in range(1, last_col + 1):
        ws.cell(total_row, col).fill = fill(NAVY)
        ws.cell(total_row, col).border = thin_border(NAVY_DEEP)
    ws.cell(total_row, 7, "без материалов").font = font(size=10, italic=True, color=GOLD_SOFT)
    ws.cell(total_row, 7).alignment = align("left", "center")
    ws.cell(total_row, 7).fill = fill(NAVY)

    row = total_row + 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    ws.cell(row, 1, rub_words(TARGET))
    ws.cell(row, 1).font = font(size=11, italic=True, color=STEEL)
    ws.cell(row, 1).alignment = align("right", "center")
    ws.cell(row, 1).fill = fill(CREAM)
    ws.row_dimensions[row].height = 20

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    ws.cell(
        row,
        1,
        "Объёмы договорные и приняты под итог 50 000 ₽. Если на объекте количество другое — строка считается как новое количество × расценка. "
        f"Дата сметы {DOC_DATE.strftime('%d.%m.%Y')}. Срок действия 30 дней.",
    )
    ws.cell(row, 1).font = font(size=9, color=MUTED)
    ws.cell(row, 1).alignment = align("left", "center")
    ws.cell(row, 1).fill = fill(CREAM)
    ws.row_dimensions[row].height = 28

    ws.print_area = f"A1:G{row}"
    return {
        "total_cell": f"F{total_row}",
        "total_row": total_row,
        "sections": section_rows,
        "total": grand_total(),
    }


def build_summary(wb: Workbook, totals: dict) -> None:
    ws = wb.create_sheet("Сводка", 0)
    apply_print(ws, landscape=True, title="Сводка")
    set_widths(ws, {"A": 4, "B": 38, "C": 14, "D": 14, "E": 4, "F": 18, "G": 16, "H": 16, "I": 16, "J": 4})
    cream_sheet(ws, 24, 10)

    paint(ws, 1, 1, 10, fill(NAVY_DEEP), height=24)
    ws.merge_cells("B1:D1")
    ws["B1"] = f"СВОДКА   ·   {OBJECT_NAME.upper()}   ·   {ROOM.upper()}   ·   ТОЛЬКО РАБОТЫ"
    ws["B1"].font = font(size=13, bold=True, color=GOLD)
    ws["B1"].alignment = align("left", "center")
    ws["B1"].fill = fill(NAVY_DEEP)
    paint(ws, 2, 1, 10, fill(GOLD), height=5)

    headers = ["Работа", "Сумма, ₽", "Доля"]
    for col, text in enumerate(headers, 2):
        cell = ws.cell(4, col, text)
        cell.font = font(size=10, bold=True, color=WHITE)
        cell.fill = fill(NAVY)
        cell.alignment = align("center")
        cell.border = thin_border(NAVY)

    items = SECTIONS[0]["items"]
    item_cells = []
    # Reconstruct item sum cells from Смета sheet layout: header 5, section 6, intro 7, items 8-10
    for i, item in enumerate(items):
        r = 5 + i
        bg = WHITE if i % 2 == 0 else ROW_ALT
        amount = round_rub(item["qty"] * item["price"])
        smeta_row = 8 + i
        ws.row_dimensions[r].height = 24
        ws.cell(r, 2, f"{item['code']}.  {item['name']}")
        ws.cell(r, 3, f"=Смета!F{smeta_row}")
        ws.cell(r, 4, f"=IF(Смета!{totals['total_cell']}=0,0,C{r}/Смета!{totals['total_cell']})")
        for col in range(2, 5):
            ws.cell(r, col).fill = fill(bg)
            ws.cell(r, col).border = thin_border(LINE_SOFT)
            ws.cell(r, col).font = font(size=11, color=INK)
        ws.cell(r, 2).alignment = align("left", "center")
        ws.cell(r, 3).alignment = align("right", "center")
        ws.cell(r, 3).number_format = "#,##0"
        ws.cell(r, 3).font = font(size=12, bold=True, color=NAVY)
        ws.cell(r, 4).alignment = align("center")
        ws.cell(r, 4).number_format = "0.0%"
        item_cells.append((item["name"], amount))

    total_r = 8
    paint(ws, total_r, 2, 4, fill(NAVY), height=28)
    ws.cell(total_r, 2, "Итого работы").font = font(size=12, bold=True, color=WHITE)
    ws.cell(total_r, 2).alignment = align("left", "center")
    ws.cell(total_r, 2).fill = fill(NAVY)
    ws.cell(total_r, 3, f"=Смета!{totals['total_cell']}")
    ws.cell(total_r, 3).font = font(size=14, bold=True, color=GOLD)
    ws.cell(total_r, 3).number_format = '#,##0" ₽"'
    ws.cell(total_r, 3).fill = fill(NAVY)
    ws.cell(total_r, 3).alignment = align("right", "center")
    ws.cell(total_r, 4, "100%").font = font(size=12, bold=True, color=GOLD_SOFT)
    ws.cell(total_r, 4).fill = fill(NAVY)
    ws.cell(total_r, 4).alignment = align("center")
    for col in range(2, 5):
        ws.cell(total_r, col).border = thin_border(NAVY_DEEP)

    ws.merge_cells("B10:D10")
    ws["B10"] = rub_words(TARGET)
    ws["B10"].font = font(size=11, italic=True, color=STEEL)
    ws["B10"].fill = fill(CREAM)
    ws.row_dimensions[10].height = 22

    short = ["Грунтовка", "Обои", "Покраска"]
    for i, name in enumerate(short):
        ws.cell(5 + i, 12, name)
    ws.column_dimensions["L"].hidden = True

    chart = BarChart()
    chart.type = "col"
    chart.title = "Работы по строкам, ₽"
    data = Reference(ws, min_col=3, min_row=4, max_row=7)
    cats = Reference(ws, min_col=12, min_row=5, max_row=7)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    chart.style = 10
    chart.y_axis.numFmt = "#,##0"
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.dataLabels.numFmt = "#,##0"
    chart.width = 14
    chart.height = 8
    chart.series[0].graphicalProperties.solidFill = NAVY
    ws.add_chart(chart, "F4")

    notes = [
        "В смету входят только три работы: грунт стен зала, поклейка обоев и покраска.",
        "Крупные строки — обои и покраска. Грунтовка — подготовка под них.",
        "Количество можно заменить: сумма строки и итог пересчитаются сами.",
    ]
    ws.merge_cells("B12:D12")
    ws["B12"] = "Как читать"
    ws["B12"].font = font(size=12, bold=True, color=NAVY)
    ws["B12"].fill = fill(CREAM)
    for i, text in enumerate(notes):
        r = 13 + i
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.cell(r, 2, f"{i + 1}.  {text}")
        ws.cell(r, 2).font = font(size=10, color=INK)
        ws.cell(r, 2).fill = fill(WHITE if i % 2 == 0 else SAND)
        ws.cell(r, 3).fill = fill(WHITE if i % 2 == 0 else SAND)
        ws.cell(r, 4).fill = fill(WHITE if i % 2 == 0 else SAND)
        ws.row_dimensions[r].height = 22

    ws.print_area = "A1:J16"


def build_cover(wb: Workbook, totals: dict) -> None:
    ws = wb.create_sheet("Титул", 0)
    apply_print(ws, landscape=False, title="Титульный лист")
    set_widths(ws, {c: 14 for c in "ABCDEFGH"})
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["H"].width = 3
    cream_sheet(ws, 44, 8)

    paint(ws, 1, 1, 8, fill(NAVY_DEEP), height=12)
    paint(ws, 2, 1, 8, fill(NAVY_DEEP), height=22)
    ws.merge_cells("B2:G2")
    ws["B2"] = "ЛОКАЛЬНАЯ СМЕТА   ·   ТОЛЬКО РАБОТЫ"
    ws["B2"].font = font(size=12, bold=True, color=GOLD)
    ws["B2"].alignment = align("left", "center")
    ws["B2"].fill = fill(NAVY_DEEP)
    paint(ws, 3, 1, 8, fill(NAVY_DEEP), height=8)
    paint(ws, 4, 1, 8, fill(GOLD), height=6)

    ws.merge_cells("B6:G6")
    ws["B6"] = "Объект Симагино"
    ws["B6"].font = font(size=28, bold=True, color=NAVY)
    ws["B6"].alignment = align("left", "center")
    ws["B6"].fill = fill(CREAM)
    ws.row_dimensions[6].height = 36

    ws.merge_cells("B7:G7")
    ws["B7"] = "Зал: грунтовка стен, поклейка обоев, покраска"
    ws["B7"].font = font(size=13, color=STEEL)
    ws["B7"].alignment = align("left", "center")
    ws["B7"].fill = fill(CREAM)
    ws.row_dimensions[7].height = 22

    ws.merge_cells("B8:G8")
    ws["B8"] = f"{DOC_NO} от {DOC_DATE.strftime('%d.%m.%Y')}   ·   материалы в смету не входят"
    ws["B8"].font = font(size=10, italic=True, color=MUTED)
    ws["B8"].fill = fill(CREAM)
    ws.row_dimensions[8].height = 18

    meta = [
        ("Объект", "Симагино"),
        ("Помещение", "Зал"),
        ("Состав", "Три работы: грунт стен, поклейка обоев, покраска"),
        ("Не входит", "Материалы, шпаклёвка, демонтаж старых обоев, вынос мебели"),
        ("Объёмы", "Договорные, приняты под фиксированный итог 50 000 ₽"),
        ("Пересчёт", "Новое количество × расценка строки"),
        ("Срок сметы", "30 календарных дней"),
        ("НДС", "Не выделен"),
    ]
    for i, (key, value) in enumerate(meta):
        r = 10 + i
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=7)
        bg = WHITE if i % 2 == 0 else SAND
        paint(ws, r, 2, 7, fill(bg), height=20, bd=thin_border(LINE_SOFT))
        ws.cell(r, 2, key).font = font(size=10, bold=True, color=NAVY)
        ws.cell(r, 2).alignment = align("left", "center")
        ws.cell(r, 4, value).font = font(size=10, color=INK)
        ws.cell(r, 4).alignment = align("left", "center")

    ws.merge_cells("B19:G19")
    paint(ws, 19, 2, 7, fill(NAVY), height=18)
    ws["B19"] = "ИТОГО РАБОТЫ"
    ws["B19"].font = font(size=10, bold=True, color=GOLD)
    ws["B19"].alignment = align("center")
    ws["B19"].fill = fill(NAVY)

    ws.merge_cells("B20:G20")
    paint(ws, 20, 2, 7, fill(NAVY), height=34)
    ws["B20"] = f"=Смета!{totals['total_cell']}"
    ws["B20"].font = font(size=28, bold=True, color=WHITE)
    ws["B20"].alignment = align("center")
    ws["B20"].number_format = '#,##0" ₽"'
    ws["B20"].fill = fill(NAVY)

    ws.merge_cells("B21:G21")
    paint(ws, 21, 2, 7, fill(NAVY), height=18)
    ws["B21"] = "без материалов"
    ws["B21"].font = font(size=9, color=GOLD_SOFT)
    ws["B21"].alignment = align("center")
    ws["B21"].fill = fill(NAVY)

    ws.merge_cells("B22:G22")
    ws["B22"] = rub_words(TARGET)
    ws["B22"].font = font(size=11, italic=True, color=STEEL)
    ws["B22"].alignment = align("center")
    ws["B22"].fill = fill(CREAM)
    ws.row_dimensions[22].height = 22

    ws.merge_cells("B24:G24")
    ws["B24"] = "Строки сметы"
    ws["B24"].font = font(size=12, bold=True, color=NAVY)
    ws["B24"].fill = fill(CREAM)

    for i, item in enumerate(SECTIONS[0]["items"]):
        r = 25 + i
        bg = WHITE if i % 2 == 0 else PALE_BLUE
        amount = round_rub(item["qty"] * item["price"])
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
        paint(ws, r, 2, 7, fill(bg), height=24, bd=thin_border(LINE_SOFT))
        ws.cell(r, 2, f"{item['code']}.  {item['name']}")
        ws.cell(r, 2).font = font(size=11, color=INK)
        ws.cell(r, 2).alignment = align("left", "center")
        ws.cell(r, 6, amount)
        ws.cell(r, 6).font = font(size=12, bold=True, color=NAVY)
        ws.cell(r, 6).number_format = '#,##0" ₽"'
        ws.cell(r, 6).alignment = align("right", "center")

    included = [
        "Грунтовка стен зала",
        "Поклейка обоев",
        "Покраска потолка и поверхностей без обоев",
    ]
    excluded = [
        "Грунт, обои, клей, краска и расходники",
        "Шпаклёвка и вышкуривание стен",
        "Демонтаж старых обоев и вынос мебели",
    ]
    ws.merge_cells("B29:D29")
    ws.merge_cells("E29:G29")
    paint(ws, 29, 2, 4, fill(GREEN), height=18)
    paint(ws, 29, 5, 7, fill(TERRACOTTA), height=18)
    ws["B29"] = "ВХОДИТ"
    ws["B29"].font = font(size=10, bold=True, color=WHITE)
    ws["B29"].alignment = align("center")
    ws["B29"].fill = fill(GREEN)
    ws["E29"] = "НЕ ВХОДИТ"
    ws["E29"].font = font(size=10, bold=True, color=WHITE)
    ws["E29"].alignment = align("center")
    ws["E29"].fill = fill(TERRACOTTA)

    for i in range(3):
        r = 30 + i
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=7)
        left_bg = WHITE if i % 2 == 0 else "F3F7F4"
        right_bg = WHITE if i % 2 == 0 else "F8F1EC"
        paint(ws, r, 2, 4, fill(left_bg), height=28)
        paint(ws, r, 5, 7, fill(right_bg), height=28)
        ws.cell(r, 2, f"{i + 1}.  {included[i]}").font = font(size=9, color=INK)
        ws.cell(r, 2).alignment = align("left", "center")
        ws.cell(r, 5, f"{i + 1}.  {excluded[i]}").font = font(size=9, color=INK)
        ws.cell(r, 5).alignment = align("left", "center")

    ws.merge_cells("B34:C34")
    ws["B34"] = "Заказчик"
    ws["B34"].font = font(size=10, bold=True, color=NAVY)
    ws["B34"].fill = fill(CREAM)
    ws.merge_cells("E34:G34")
    ws["E34"] = "Исполнитель"
    ws["E34"].font = font(size=10, bold=True, color=NAVY)
    ws["E34"].fill = fill(CREAM)

    ws.merge_cells("B35:C35")
    ws["B35"] = "______________________________"
    ws["B35"].font = font(size=10, color=NAVY)
    ws["B35"].fill = fill(CREAM)
    ws.merge_cells("E35:G35")
    ws["E35"] = "______________________________"
    ws["E35"].font = font(size=10, color=NAVY)
    ws["E35"].fill = fill(CREAM)

    ws.merge_cells("B36:C36")
    ws["B36"] = "подпись, расшифровка"
    ws["B36"].font = font(size=8, italic=True, color=MUTED)
    ws["B36"].fill = fill(CREAM)
    ws.merge_cells("E36:G36")
    ws["E36"] = "подпись, расшифровка"
    ws["E36"].font = font(size=8, italic=True, color=MUTED)
    ws["E36"].fill = fill(CREAM)

    ws.merge_cells("B38:G38")
    ws["B38"] = (
        "Смета по объекту Симагино, помещение зал. Итог работ 50 000 ₽. "
        "Объёмы договорные — лист «Объёмы»."
    )
    ws["B38"].font = font(size=8, italic=True, color=MUTED)
    ws["B38"].alignment = align("left", "center")
    ws["B38"].fill = fill(CREAM)
    ws.row_dimensions[38].height = 28

    ws.print_area = "A1:H39"
    ws.page_setup.fitToHeight = 1


def build_volumes(wb: Workbook) -> None:
    ws = wb.create_sheet("Объёмы")
    apply_print(ws, landscape=True, title="Договорные объёмы")
    set_widths(ws, {"A": 6, "B": 48, "C": 16, "D": 14, "E": 48})
    cream_sheet(ws, 16, 5)

    paint(ws, 1, 1, 5, fill(NAVY_DEEP), height=24)
    ws.merge_cells("A1:E1")
    ws["A1"] = "ДОГОВОРНЫЕ ОБЪЁМЫ   ·   ЗАЛ   ·   ПОД ИТОГ 50 000 ₽"
    ws["A1"].font = font(size=13, bold=True, color=GOLD)
    ws["A1"].alignment = align("left", "center")
    ws["A1"].fill = fill(NAVY_DEEP)
    paint(ws, 2, 1, 5, fill(GOLD), height=5)

    ws.merge_cells("A3:E3")
    ws["A3"] = "Количество ниже — основание цены. Если на объекте иначе: новое количество × расценка листа «Смета»."
    ws["A3"].font = font(size=10, color=STEEL)
    ws["A3"].fill = fill(CREAM)
    ws.row_dimensions[3].height = 28

    for col, text in enumerate(["№", "Показатель", "Объём", "Пункт", "Пояснение"], 1):
        cell = ws.cell(5, col, text)
        cell.font = font(size=10, bold=True, color=WHITE)
        cell.fill = fill(NAVY)
        cell.alignment = align("center")
        cell.border = thin_border(NAVY)

    for i, (name, qty, point, note) in enumerate(VOLUMES):
        r = 6 + i
        bg = WHITE if i % 2 == 0 else ROW_ALT
        ws.row_dimensions[r].height = 24
        for col, (value, how) in enumerate(
            zip([i + 1, name, qty, point, note], ["center", "left", "center", "center", "left"]),
            1,
        ):
            cell = ws.cell(r, col, value)
            cell.font = font(size=10, bold=(col == 3), color=INK)
            cell.fill = fill(bg)
            cell.alignment = align(how)
            cell.border = thin_border(LINE_SOFT)

    ws.print_area = "A1:E8"


def build_conditions(wb: Workbook) -> None:
    ws = wb.create_sheet("Условия")
    apply_print(ws, landscape=False, title="Условия")
    set_widths(ws, {"A": 4, "B": 8, "C": 78, "D": 4})
    cream_sheet(ws, 28, 4)

    paint(ws, 1, 1, 4, fill(NAVY_DEEP), height=24)
    ws.merge_cells("B1:C1")
    ws["B1"] = "УСЛОВИЯ СМЕТЫ"
    ws["B1"].font = font(size=14, bold=True, color=GOLD)
    ws["B1"].alignment = align("left", "center")
    ws["B1"].fill = fill(NAVY_DEEP)
    paint(ws, 2, 1, 4, fill(GOLD), height=5)

    blocks = [
        (
            "Цена",
            [
                "Итог 50 000 ₽ — стоимость трёх работ по залу.",
                "Материалы, доставка и подъём в эту сумму не входят.",
                "НДС не выделен.",
            ],
        ),
        (
            "Объёмы",
            [
                "Грунтовка стен зала — 50 м².",
                "Поклейка обоев — 50 м².",
                "Покраска — 40 м² (потолок и поверхности без обоев, два слоя).",
            ],
        ),
        (
            "Границы работ",
            [
                "Шпаклёвка и шлифовка стен в эту смету не входят.",
                "Демонтаж старых обоев не входит.",
                "Вынос и занос мебели не входит.",
                "Материалы закупает заказчик.",
            ],
        ),
        (
            "Порядок",
            [
                "Срок действия расценок — 30 дней с даты сметы.",
                "Оплата и график — в договоре, не в этой смете.",
            ],
        ),
    ]
    row = 4
    number = 1
    for title, lines in blocks:
        paint(ws, row, 2, 3, fill(STEEL), height=20)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        ws.cell(row, 2, title).font = font(size=11, bold=True, color=WHITE)
        ws.cell(row, 2).alignment = align("left", "center")
        ws.cell(row, 2).fill = fill(STEEL)
        ws.cell(row, 3).fill = fill(STEEL)
        row += 1
        for line in lines:
            bg = WHITE if number % 2 else SAND
            paint(ws, row, 2, 3, fill(bg), height=24, bd=thin_border(LINE_SOFT))
            ws.cell(row, 2, f"{number}.").font = font(size=10, bold=True, color=NAVY)
            ws.cell(row, 2).alignment = align("center")
            ws.cell(row, 3, line).font = font(size=10, color=INK)
            ws.cell(row, 3).alignment = align("left", "center")
            number += 1
            row += 1
        row += 1

    ws.print_area = f"A1:D{row}"


def build_materials(wb: Workbook) -> None:
    ws = wb.create_sheet("Материалы заказчика")
    apply_print(ws, landscape=False, title="Материалы заказчика")
    set_widths(ws, {"A": 6, "B": 18, "C": 78})
    cream_sheet(ws, 16, 3)

    paint(ws, 1, 1, 3, fill(NAVY_DEEP), height=24)
    ws.merge_cells("A1:C1")
    ws["A1"] = "МАТЕРИАЛЫ ЗАКАЗЧИКА   ·   В 50 000 ₽ НЕ ВХОДЯТ"
    ws["A1"].font = font(size=13, bold=True, color=GOLD)
    ws["A1"].alignment = align("left", "center")
    ws["A1"].fill = fill(NAVY_DEEP)
    paint(ws, 2, 1, 3, fill(GOLD), height=5)

    ws.merge_cells("A3:C3")
    ws["A3"] = "Список для закупки. Цен материалов в смете нет."
    ws["A3"].font = font(size=10, italic=True, color=MUTED)
    ws["A3"].fill = fill(CREAM)

    for col, text in enumerate(["№", "Группа", "Что подготовить"], 1):
        cell = ws.cell(5, col, text)
        cell.font = font(size=10, bold=True, color=WHITE)
        cell.fill = fill(NAVY)
        cell.alignment = align("center")
        cell.border = thin_border(NAVY)

    for i, (group, item) in enumerate(MATERIALS):
        r = 6 + i
        bg = WHITE if i % 2 == 0 else ROW_ALT
        for col, value in enumerate([i + 1, group, item], 1):
            cell = ws.cell(r, col, value)
            cell.font = font(size=10, bold=(col == 2), color=INK)
            cell.fill = fill(bg)
            cell.alignment = align("center" if col < 3 else "left")
            cell.border = thin_border(LINE_SOFT)
        ws.row_dimensions[r].height = 22

    last = 6 + len(MATERIALS)
    ws.print_area = f"A1:C{last}"


def assert_total() -> None:
    total = grand_total()
    if total != TARGET:
        raise SystemExit(f"Итог {total} != {TARGET}")


def verify_workbook(path: Path) -> None:
    book = load_workbook(path)
    sheet = book["Смета"]
    line_total = 0
    codes = []
    for row in range(1, sheet.max_row + 1):
        code = sheet.cell(row, 1).value
        if not (isinstance(code, str) and len(code) >= 3 and code[0].isdigit() and "." in code):
            continue
        qty = sheet.cell(row, 4).value
        price = sheet.cell(row, 5).value
        formula = sheet.cell(row, 6).value
        if formula != f"=ROUND(D{row}*E{row},0)":
            raise SystemExit(f"Строка {code}: формула {formula}")
        line_total += round_rub(qty * price)
        codes.append(code)
    expected = [item["code"] for section in SECTIONS for item in section["items"]]
    if codes != expected:
        raise SystemExit(f"Состав строк не совпал: {codes}")
    if line_total != TARGET:
        raise SystemExit(f"Сумма строк {line_total} != {TARGET}")
    if book["Титул"]["B6"].value != "Объект Симагино":
        raise SystemExit("Титул не про Симагино")
    if book["Титул"]["B7"].value != "Зал: грунтовка стен, поклейка обоев, покраска":
        raise SystemExit("Титул не про зал")
    if rub_words(TARGET) != "Пятьдесят тысяч рублей 00 копеек":
        raise SystemExit(rub_words(TARGET))
    print("CHECK OK", line_total, "lines", len(codes))


def main() -> None:
    assert_total()
    wb = Workbook()
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True)
    totals = build_estimate(wb)
    build_summary(wb, totals)
    build_cover(wb, totals)
    build_volumes(wb)
    build_conditions(wb)
    build_materials(wb)

    order = ["Титул", "Сводка", "Смета", "Объёмы", "Условия", "Материалы заказчика"]
    for index, name in enumerate(order):
        current = wb.sheetnames.index(name)
        wb.move_sheet(name, offset=index - current)

    colors = {
        "Титул": GOLD,
        "Сводка": STEEL,
        "Смета": NAVY,
        "Объёмы": GREEN,
        "Условия": TERRACOTTA,
        "Материалы заказчика": GOLD_SOFT,
    }
    for name, color in colors.items():
        wb[name].sheet_properties.tabColor = color

    wb.properties.title = "Смета работ · Симагино · Зал"
    wb.properties.subject = "Только работы, 50 000 руб."
    wb.properties.creator = "Локальная смета"
    wb.properties.description = "Объект Симагино, зал: грунтовка стен, поклейка обоев, покраска. Только работы."
    wb.active = wb["Титул"]
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)

    print("FILE", OUT)
    print("TOTAL", totals["total"])
    print("WORDS", rub_words(totals["total"]))
    for item in SECTIONS[0]["items"]:
        print(f"ITEM {item['code']} {round_rub(item['qty'] * item['price'])} {item['name']}")
    verify_workbook(OUT)


if __name__ == "__main__":
    main()
