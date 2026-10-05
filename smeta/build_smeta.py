#!/usr/bin/env python3
"""Локальная смета на работы. Объект: Симагино. Только работы, без материалов.

Итог по перечню заказчика зафиксирован: 230 000 ₽.
Объёмы договорные — под эту сумму, расценки можно пересчитать при замере.
"""

from __future__ import annotations

from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.page import PageMargins


NAVY = "1B365D"
NAVY_DEEP = "12243F"
STEEL = "3D5A80"
GOLD = "C4A35A"
GOLD_SOFT = "E6D3A3"
CREAM = "FBF8F2"
SAND = "F3EBDD"
ROW_ALT = "F7F3EA"
WHITE = "FFFFFF"
INK = "1F2933"
MUTED = "5C6670"
LINE = "D7CDBD"
LINE_SOFT = "E8E0D4"
SUBTOTAL_BG = "E7EEF6"
TOTAL_BG = "1B365D"
GREEN = "2F6B4F"
TERRACOTTA = "A65D3F"
PALE_BLUE = "EEF3F8"

TARGET = 230_000
DOC_NO = "ЛС-СИМ-01"
OBJECT_NAME = "Симагино"
DOC_DATE = date(2026, 10, 5)
OUT = Path(__file__).resolve().parent / "Smeta_raboty_Simagino.xlsx"


def round_rub(n: float) -> int:
    return int(Decimal(str(n)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def money(n: float) -> str:
    return f"{int(round_rub(n)):,.0f}".replace(",", " ")


def rub_words(n: int) -> str:
    """Сумма прописью, рубли, без копеек в числе — «00 копеек»."""
    units = (
        ("", "", ""),
        ("тысяча", "тысячи", "тысяч"),
        ("миллион", "миллиона", "миллионов"),
    )
    ones_m = ["", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
    ones_f = ["", "одна", "две", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
    teens = [
        "десять", "одиннадцать", "двенадцать", "тринадцать", "четырнадцать",
        "пятнадцать", "шестнадцать", "семнадцать", "восемнадцать", "девятнадцать",
    ]
    tens = ["", "", "двадцать", "тридцать", "сорок", "пятьдесят", "шестьдесят", "семьдесят", "восемьдесят", "девяносто"]
    hundreds = ["", "сто", "двести", "триста", "четыреста", "пятьсот", "шестьсот", "семьсот", "восемьсот", "девятьсот"]

    def tri(num: int, feminine: bool) -> str:
        h, rest = divmod(num, 100)
        t, o = divmod(rest, 10)
        ones = ones_f if feminine else ones_m
        parts = []
        if h:
            parts.append(hundreds[h])
        if t == 1:
            parts.append(teens[o])
        else:
            if t:
                parts.append(tens[t])
            if o:
                parts.append(ones[o])
        return " ".join(parts)

    def form(num: int, forms: tuple[str, str, str]) -> str:
        n100 = num % 100
        n10 = num % 10
        if 11 <= n100 <= 19:
            return forms[2]
        if n10 == 1:
            return forms[0]
        if 2 <= n10 <= 4:
            return forms[1]
        return forms[2]

    n = int(n)
    if n == 0:
        return "Ноль рублей 00 копеек"
    chunks = []
    tmp = n
    while tmp:
        chunks.append(tmp % 1000)
        tmp //= 1000
    words = []
    for i, chunk in enumerate(chunks):
        if chunk == 0:
            continue
        w = tri(chunk, feminine=(i == 1))
        if i:
            w = f"{w} {form(chunk, units[i])}".strip()
        words.append(w)
    words.reverse()
    body = " ".join(words).strip()
    body = body[0].upper() + body[1:]
    return f"{body} {form(n, ('рубль', 'рубля', 'рублей'))} 00 копеек"


# kind of scope the customer named, nothing else.
SECTIONS: list[dict] = [
    {
        "num": "1",
        "title": "Электромонтажные работы",
        "intro": "Завершение электрики по уже намеченным трассам. Новый щит не собирается.",
        "items": [
            {
                "code": "1.1",
                "name": "Вмазка коробок",
                "unit": "шт.",
                "qty": 40,
                "price": 400,
                "note": "Подрозетники и установочные коробки. Смесь — заказчика",
            },
            {
                "code": "1.2",
                "name": "Штробление и прокладка проводов",
                "unit": "м.п.",
                "qty": 60,
                "price": 800,
                "note": "Резка штрабы и укладка провода. Заделка штрабы не входит",
            },
            {
                "code": "1.3",
                "name": "Демонтаж 2-го щита",
                "unit": "шт.",
                "qty": 1,
                "price": 5_000,
                "note": "Снятие щита и отключение его линий",
            },
            {
                "code": "1.4",
                "name": "Прозвонка проводов",
                "unit": "линия",
                "qty": 40,
                "price": 150,
                "note": "Прозвонка и маркировка линий",
            },
            {
                "code": "1.5",
                "name": "Распайка коробок",
                "unit": "шт.",
                "qty": 20,
                "price": 700,
                "note": "Соединения в распаечных коробках",
            },
        ],
    },
    {
        "num": "2",
        "title": "Подготовка стен, 1 этаж",
        "intro": "Только стены первого этажа. Потолок и второй этаж не входят. Окраска и обои не входят.",
        "items": [
            {
                "code": "2.1",
                "name": "Грунтовка стен 1 этажа",
                "unit": "м²",
                "qty": 120,
                "price": 75,
                "note": "Один слой. Площадь стен за вычетом проёмов",
            },
            {
                "code": "2.2",
                "name": "Шпаклёвка, 2 слоя",
                "unit": "м²",
                "qty": 120,
                "price": 450,
                "note": "Два слоя под обои или окраску",
            },
            {
                "code": "2.3",
                "name": "Вышкуривание",
                "unit": "м²",
                "qty": 120,
                "price": 100,
                "note": "Шлифовка шпаклёвки. Абразив — заказчика",
            },
        ],
    },
    {
        "num": "3",
        "title": "Откосы, подоконники, короб над печкой",
        "intro": "Изделия и листовой материал — заказчика. Чистовая окраска этих узлов не входит.",
        "items": [
            {
                "code": "3.1",
                "name": "Устройство откосов в дверном блоке",
                "unit": "м.п.",
                "qty": 6,
                "price": 1_500,
                "note": "Один дверной блок: стойки и вершник",
            },
            {
                "code": "3.2",
                "name": "Монтаж дубовых подоконников",
                "unit": "шт.",
                "qty": 4,
                "price": 2_500,
                "note": "Подгонка и крепление. Масло и лак не входят",
            },
            {
                "code": "3.3",
                "name": "Устройство короба над печкой",
                "unit": "компл.",
                "qty": 1,
                "price": 11_000,
                "note": "Каркас и обшивка. Чистовая отделка короба не входит",
            },
        ],
    },
    {
        "num": "4",
        "title": "Котельная: трубы, коллекторы, бойлер",
        "intro": "Гидравлическая часть. Котёл не монтируется и не запускается. Бойлер электрически не подключается.",
        "items": [
            {
                "code": "4.1",
                "name": "Укладка труб до котла",
                "unit": "м.п.",
                "qty": 15,
                "price": 800,
                "note": "Прокладка и крепление трассы. Пуск котла не входит",
            },
            {
                "code": "4.2",
                "name": "Монтаж коллекторов, фильтра и подключение труб",
                "unit": "компл.",
                "qty": 1,
                "price": 15_000,
                "note": "Сборка узла и подключение труб. Оборудование — заказчика",
            },
            {
                "code": "4.3",
                "name": "Обвязка бойлера",
                "unit": "компл.",
                "qty": 1,
                "price": 9_000,
                "note": "Обвязка по воде. Электроподключение не входит",
            },
        ],
    },
]

VOLUMES = [
    ("Стены 1 этажа под грунт, шпаклёвку и шлифовку", "120 м²", "п. 2.1–2.3", "За вычетом проёмов. Потолок не считается"),
    ("Штрабы с укладкой провода", "60 м.п.", "п. 1.2", "Штробление и прокладка одной строкой"),
    ("Коробки под вмазку", "40 шт.", "п. 1.1", "Подрозетники и установочные коробки"),
    ("Линии на прозвонку", "40 линий", "п. 1.4", "С маркировкой"),
    ("Распаечные коробки", "20 шт.", "п. 1.5", "Соединения внутри коробки"),
    ("Второй электрический щит", "1 шт.", "п. 1.3", "Только демонтаж"),
    ("Откосы дверного блока", "6 м.п.", "п. 3.1", "Один блок, стойки и вершник"),
    ("Дубовые подоконники", "4 шт.", "п. 3.2", "Изделия заказчика"),
    ("Короб над печкой", "1 компл.", "п. 3.3", "Каркас и обшивка"),
    ("Трасса труб до котла", "15 м.п.", "п. 4.1", "Без установки котла"),
    ("Коллекторы и фильтр", "1 компл.", "п. 4.2", "Монтаж узла и подключение труб"),
    ("Бойлер", "1 компл.", "п. 4.3", "Гидравлическая обвязка"),
]

MATERIALS = [
    ("Электрика", "Подрозетники, распаечные коробки, алебастр или гипс"),
    ("Электрика", "Кабель и провод, клеммы для распайки"),
    ("Стены", "Грунт"),
    ("Стены", "Шпаклёвка на два слоя"),
    ("Стены", "Абразивная сетка и наждачная бумага"),
    ("Откосы", "Смесь или листы, профиль и уголки на один дверной блок"),
    ("Подоконники", "Дубовые подоконники, 4 шт., крепёж и пена"),
    ("Печка", "Брусок или профиль и обшивка короба над печкой"),
    ("Котельная", "Труба, фитинги и крепёж трассы до котла"),
    ("Котельная", "Коллекторы, фильтр, краны и фитинги узла"),
    ("Котельная", "Комплект обвязки бойлера: краны, группа безопасности, американки"),
]


def section_sum(section: dict) -> int:
    return sum(round_rub(item["qty"] * item["price"]) for item in section["items"])


def grand_total() -> int:
    return sum(section_sum(section) for section in SECTIONS)


def thin_border(color: str = LINE) -> Border:
    side = Side(style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def font(name: str = "Calibri", size: int = 11, bold: bool = False, color: str = INK, italic: bool = False) -> Font:
    return Font(name=name, size=size, bold=bold, color=color, italic=italic)


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def align(h: str = "left", v: str = "center", wrap: bool = True) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def apply_print(ws, landscape: bool, title: str) -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.45, right=0.45, top=0.65, bottom=0.55, header=0.25, footer=0.25)
    ws.oddHeader.left.text = f"&K1B365D{DOC_NO}  ·  {OBJECT_NAME}"
    ws.oddHeader.right.text = "&K5C6670только работы"
    ws.oddFooter.left.text = f"&K5C6670{title}"
    ws.oddFooter.center.text = "&K5C6670стр. &P из &N"
    ws.oddFooter.right.text = "&K5C6670230 000 ₽"
    ws.print_options.horizontalCentered = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.view = "normal"
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def set_widths(ws, widths: dict[str, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def paint(ws, row: int, c1: int, c2: int, fl=None, ft=None, al=None, bd=None, height: float | None = None) -> None:
    if height is not None:
        ws.row_dimensions[row].height = height
    for col in range(c1, c2 + 1):
        cell = ws.cell(row, col)
        if fl:
            cell.fill = fl
        if ft:
            cell.font = ft
        if al:
            cell.alignment = al
        if bd:
            cell.border = bd


def cream_sheet(ws, rows: int, cols: int) -> None:
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            ws.cell(r, c).fill = fill(CREAM)


def build_estimate(wb: Workbook) -> dict:
    ws = wb.active
    ws.title = "Смета"
    apply_print(ws, landscape=True, title="Локальная смета")
    set_widths(ws, {"A": 8, "B": 58, "C": 12, "D": 12, "E": 14, "F": 16, "G": 48})
    last_col = 7
    cream_sheet(ws, 48, last_col)

    paint(ws, 1, 1, last_col, fill(NAVY_DEEP), height=22)
    ws.merge_cells("A1:G1")
    ws["A1"] = f"ЛОКАЛЬНАЯ СМЕТА НА РАБОТЫ   ·   {DOC_NO}   ·   ОБЪЕКТ {OBJECT_NAME.upper()}"
    ws["A1"].font = font(size=12, bold=True, color=GOLD)
    ws["A1"].alignment = align("left", "center")
    ws["A1"].fill = fill(NAVY_DEEP)

    paint(ws, 2, 1, last_col, fill(GOLD), height=5)

    ws.merge_cells("A3:G3")
    ws["A3"] = "Только работы. Материалы, оборудование и крепёж в итог не входят."
    ws["A3"].font = font(size=10, italic=True, color=MUTED)
    ws["A3"].alignment = align("left", "center")
    ws["A3"].fill = fill(CREAM)
    ws.row_dimensions[3].height = 18

    headers = ["№", "Наименование работ", "Ед. изм.", "Кол-во", "Расценка, ₽", "Сумма, ₽", "Что входит в строку"]
    header_row = 5
    ws.row_dimensions[header_row].height = 24
    for col, text in enumerate(headers, 1):
        cell = ws.cell(header_row, col, text)
        cell.font = font(size=10, bold=True, color=WHITE)
        cell.fill = fill(NAVY)
        cell.alignment = align("center", "center")
        cell.border = thin_border(NAVY)
    ws.freeze_panes = "A6"
    ws.print_title_rows = "1:5"
    ws.page_setup.fitToHeight = 1

    row = 6
    section_rows = []
    money_format = '#,##0'
    for section in SECTIONS:
        paint(ws, row, 1, last_col, fill(STEEL), height=22)
        ws.cell(row, 1, section["num"]).font = font(size=12, bold=True, color=GOLD)
        ws.cell(row, 1).alignment = align("center")
        ws.cell(row, 1).fill = fill(STEEL)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col)
        ws.cell(row, 2, section["title"]).font = font(size=12, bold=True, color=WHITE)
        ws.cell(row, 2).alignment = align("left", "center")
        ws.cell(row, 2).fill = fill(STEEL)
        for col in range(1, last_col + 1):
            ws.cell(row, col).fill = fill(STEEL)
            ws.cell(row, col).border = thin_border(STEEL)
        row += 1

        paint(ws, row, 1, last_col, fill(SAND), height=20)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
        ws.cell(row, 1, section["intro"]).font = font(size=9, italic=True, color=STEEL)
        ws.cell(row, 1).alignment = align("left", "center")
        ws.cell(row, 1).fill = fill(SAND)
        row += 1

        first_item = row
        for item_i, item in enumerate(section["items"]):
            bg = WHITE if item_i % 2 == 0 else ROW_ALT
            ws.row_dimensions[row].height = 28
            values = [
                item["code"],
                item["name"],
                item["unit"],
                item["qty"],
                item["price"],
                f"=ROUND(D{row}*E{row},0)",
                item["note"],
            ]
            aligns = ["center", "left", "center", "center", "right", "right", "left"]
            for col, (value, how) in enumerate(zip(values, aligns), 1):
                cell = ws.cell(row, col, value)
                cell.font = font(size=11, bold=(col == 6), color=INK)
                cell.fill = fill(bg)
                cell.alignment = align(how, "center")
                cell.border = thin_border(LINE_SOFT)
            ws.cell(row, 2).font = font(size=11, bold=True, color=INK)
            ws.cell(row, 4).number_format = "0"
            ws.cell(row, 5).number_format = money_format
            ws.cell(row, 6).number_format = money_format
            row += 1
        last_item = row - 1

        paint(ws, row, 1, last_col, fill(SUBTOTAL_BG), height=22)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        ws.cell(row, 1, f"Итого по разделу {section['num']}.  {section['title']}")
        ws.cell(row, 1).font = font(size=10, bold=True, color=NAVY)
        ws.cell(row, 1).alignment = align("right", "center")
        ws.cell(row, 1).fill = fill(SUBTOTAL_BG)
        for col in range(1, 6):
            ws.cell(row, col).fill = fill(SUBTOTAL_BG)
            ws.cell(row, col).border = thin_border(LINE)
            ws.cell(row, col).font = font(size=10, bold=True, color=NAVY)
        sum_cell = ws.cell(row, 6, f"=SUM(F{first_item}:F{last_item})")
        sum_cell.font = font(size=12, bold=True, color=NAVY)
        sum_cell.fill = fill(SUBTOTAL_BG)
        sum_cell.alignment = align("right", "center")
        sum_cell.number_format = money_format
        sum_cell.border = thin_border(LINE)
        ws.cell(row, 7).fill = fill(SUBTOTAL_BG)
        ws.cell(row, 7).border = thin_border(LINE)
        section_rows.append(
            {
                "num": section["num"],
                "title": section["title"],
                "sum": section_sum(section),
                "cell": f"F{row}",
                "row": row,
            }
        )
        row += 1

    row += 1
    total_row = row
    paint(ws, total_row, 1, last_col, fill(TOTAL_BG), height=30)
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=5)
    ws.cell(total_row, 1, "ИТОГО РАБОТЫ")
    ws.cell(total_row, 1).font = font(size=14, bold=True, color=WHITE)
    ws.cell(total_row, 1).alignment = align("right", "center")
    parts = "+".join(section["cell"] for section in section_rows)
    total_cell = ws.cell(total_row, 6, f"={parts}")
    total_cell.font = font(size=16, bold=True, color=GOLD)
    total_cell.alignment = align("right", "center")
    total_cell.number_format = '#,##0" ₽"'
    for col in range(1, last_col + 1):
        ws.cell(total_row, col).fill = fill(TOTAL_BG)
        ws.cell(total_row, col).border = thin_border(NAVY_DEEP)
        if col != 6 and col != 1:
            ws.cell(total_row, col).font = font(size=12, bold=True, color=WHITE)
    ws.cell(total_row, 7, "без материалов").font = font(size=10, italic=True, color=GOLD_SOFT)
    ws.cell(total_row, 7).alignment = align("left", "center")
    ws.cell(total_row, 7).fill = fill(TOTAL_BG)

    row = total_row + 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    ws.cell(row, 1, rub_words(TARGET))
    ws.cell(row, 1).font = font(size=11, italic=True, color=STEEL)
    ws.cell(row, 1).alignment = align("right", "center")
    ws.cell(row, 1).fill = fill(CREAM)
    ws.row_dimensions[row].height = 20

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    ws.cell(
        row,
        1,
        "Объёмы договорные и приняты под итог 230 000 ₽. Если на объекте количество другое — строка считается как новое количество × расценка. "
        f"Дата сметы {DOC_DATE.strftime('%d.%m.%Y')}. Срок действия 30 дней.",
    )
    ws.cell(row, 1).font = font(size=9, color=MUTED)
    ws.cell(row, 1).alignment = align("left", "center")
    ws.cell(row, 1).fill = fill(CREAM)
    ws.row_dimensions[row].height = 28

    ws.print_area = f"A1:G{row}"
    ws.page_setup.fitToHeight = 1
    ws.sheet_view.showGridLines = False
    return {
        "total_cell": f"F{total_row}",
        "total_row": total_row,
        "sections": section_rows,
        "total": grand_total(),
    }


def build_summary(wb: Workbook, totals: dict) -> None:
    ws = wb.create_sheet("Сводка", 0)
    apply_print(ws, landscape=True, title="Сводка")
    set_widths(ws, {"A": 4, "B": 42, "C": 18, "D": 14, "E": 4, "F": 22, "G": 18, "H": 18, "I": 18, "J": 4})
    cream_sheet(ws, 28, 10)

    paint(ws, 1, 1, 10, fill(NAVY_DEEP), height=24)
    ws.merge_cells("B1:D1")
    ws["B1"] = f"СВОДКА   ·   {OBJECT_NAME.upper()}   ·   ТОЛЬКО РАБОТЫ"
    ws["B1"].font = font(size=14, bold=True, color=GOLD)
    ws["B1"].alignment = align("left", "center")
    ws["B1"].fill = fill(NAVY_DEEP)

    paint(ws, 2, 1, 10, fill(GOLD), height=5)

    headers = ["Раздел", "Сумма, ₽", "Доля"]
    for col, text in enumerate(headers, 2):
        cell = ws.cell(4, col, text)
        cell.font = font(size=10, bold=True, color=WHITE)
        cell.fill = fill(NAVY)
        cell.alignment = align("center")
        cell.border = thin_border(NAVY)
    ws.row_dimensions[4].height = 22

    for i, section in enumerate(totals["sections"]):
        r = 5 + i
        bg = WHITE if i % 2 == 0 else ROW_ALT
        ws.row_dimensions[r].height = 24
        ws.cell(r, 2, f"{section['num']}.  {section['title']}")
        ws.cell(r, 3, f"=Смета!{section['cell']}")
        ws.cell(r, 4, f"=IF(Смета!{totals['total_cell']}=0,0,C{r}/Смета!{totals['total_cell']})")
        for col in range(2, 5):
            ws.cell(r, col).fill = fill(bg)
            ws.cell(r, col).font = font(size=11, color=INK)
            ws.cell(r, col).border = thin_border(LINE_SOFT)
            ws.cell(r, col).alignment = align("left" if col == 2 else "right" if col == 3 else "center")
        ws.cell(r, 3).number_format = '#,##0'
        ws.cell(r, 3).font = font(size=12, bold=True, color=NAVY)
        ws.cell(r, 4).number_format = "0.0%"

    total_r = 9
    paint(ws, total_r, 2, 4, fill(NAVY), height=28)
    ws.cell(total_r, 2, "Итого работы").font = font(size=12, bold=True, color=WHITE)
    ws.cell(total_r, 2).alignment = align("left", "center")
    ws.cell(total_r, 2).fill = fill(NAVY)
    ws.cell(total_r, 3, f"=Смета!{totals['total_cell']}")
    ws.cell(total_r, 3).font = font(size=14, bold=True, color=GOLD)
    ws.cell(total_r, 3).number_format = '#,##0" ₽"'
    ws.cell(total_r, 3).fill = fill(NAVY)
    ws.cell(total_r, 3).alignment = align("right", "center")
    ws.cell(total_r, 4, "100%").font = font(size=12, bold=True, color=GOLD_SOFT)
    ws.cell(total_r, 4).fill = fill(NAVY)
    ws.cell(total_r, 4).alignment = align("center")
    for col in range(2, 5):
        ws.cell(total_r, col).border = thin_border(NAVY_DEEP)

    ws.merge_cells("B11:D11")
    ws["B11"] = rub_words(TARGET)
    ws["B11"].font = font(size=11, italic=True, color=STEEL)
    ws["B11"].alignment = align("left", "center")
    ws["B11"].fill = fill(CREAM)
    ws.row_dimensions[11].height = 22

    ws.merge_cells("B13:D13")
    ws["B13"] = "Как читать сводку"
    ws["B13"].font = font(size=12, bold=True, color=NAVY)
    ws["B13"].fill = fill(CREAM)

    notes = [
        "В итог входят только 14 работ из задания. Ни одна посторонняя работа не добавлена.",
        "Крупные строки — штробление с прокладкой провода и шпаклёвка стен в два слоя.",
        "Электромонтаж и подготовка стен вместе дают большую часть суммы.",
        "Котельная — трубы до котла, коллекторный узел с фильтром и обвязка бойлера.",
        "Количество в смете можно заменить: сумма строки и итог пересчитаются сами.",
    ]
    for i, text in enumerate(notes):
        r = 14 + i
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.cell(r, 2, f"{i + 1}.  {text}")
        ws.cell(r, 2).font = font(size=10, color=INK)
        ws.cell(r, 2).alignment = align("left", "center")
        ws.cell(r, 2).fill = fill(WHITE if i % 2 == 0 else SAND)
        ws.cell(r, 3).fill = fill(WHITE if i % 2 == 0 else SAND)
        ws.cell(r, 4).fill = fill(WHITE if i % 2 == 0 else SAND)
        ws.row_dimensions[r].height = 22

    short_names = ["Электрика", "Стены 1 этажа", "Откосы и короб", "Котельная"]
    for i, short in enumerate(short_names):
        ws.cell(5 + i, 12, short)
    ws.column_dimensions["L"].hidden = True

    chart = BarChart()
    chart.type = "col"
    chart.title = "Работы по разделам, ₽"
    chart.y_axis.title = None
    chart.x_axis.title = None
    data = Reference(ws, min_col=3, min_row=4, max_row=8)
    cats = Reference(ws, min_col=12, min_row=5, max_row=8)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend = None
    chart.style = 10
    chart.y_axis.numFmt = '#,##0'
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.dataLabels.numFmt = '#,##0'
    chart.width = 15
    chart.height = 8
    series = chart.series[0]
    series.graphicalProperties.solidFill = NAVY
    ws.add_chart(chart, "F4")

    ws.print_area = "A1:J20"
    ws.page_setup.fitToHeight = 1


def build_cover(wb: Workbook, totals: dict) -> None:
    ws = wb.create_sheet("Титул", 0)
    apply_print(ws, landscape=False, title="Титульный лист")
    set_widths(ws, {c: 14 for c in "ABCDEFGH"})
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["H"].width = 3
    cream_sheet(ws, 52, 8)

    paint(ws, 1, 1, 8, fill(NAVY_DEEP), height=12)
    paint(ws, 2, 1, 8, fill(NAVY_DEEP), height=22)
    ws.merge_cells("B2:G2")
    ws["B2"] = "ЛОКАЛЬНАЯ СМЕТА   ·   ТОЛЬКО РАБОТЫ"
    ws["B2"].font = font(size=12, bold=True, color=GOLD)
    ws["B2"].alignment = align("left", "center")
    ws["B2"].fill = fill(NAVY_DEEP)
    paint(ws, 3, 1, 8, fill(NAVY_DEEP), height=8)
    paint(ws, 4, 1, 8, fill(GOLD), height=6)

    ws.merge_cells("B6:G6")
    ws["B6"] = "Объект Симагино"
    ws["B6"].font = font(size=28, bold=True, color=NAVY)
    ws["B6"].alignment = align("left", "center")
    ws["B6"].fill = fill(CREAM)
    ws.row_dimensions[6].height = 36

    ws.merge_cells("B7:G7")
    ws["B7"] = "Электрика, стены 1 этажа, откосы, подоконники, короб над печкой, котельная"
    ws["B7"].font = font(size=12, color=STEEL)
    ws["B7"].alignment = align("left", "center")
    ws["B7"].fill = fill(CREAM)
    ws.row_dimensions[7].height = 20

    ws.merge_cells("B8:G8")
    ws["B8"] = f"{DOC_NO} от {DOC_DATE.strftime('%d.%m.%Y')}   ·   материалы в смету не входят"
    ws["B8"].font = font(size=10, italic=True, color=MUTED)
    ws["B8"].fill = fill(CREAM)
    ws.row_dimensions[8].height = 18

    meta = [
        ("Объект", "Симагино"),
        ("Состав", "14 работ из перечня. Только работа, без материалов и оборудования"),
        ("Не входит", "Заделка штраб, покраска, обои, лак подоконников, монтаж и пуск котла"),
        ("Не входит", "Электроподключение бойлера, новый щит, потолки, 2 этаж, вывоз мусора"),
        ("Объёмы", "Договорные, приняты под фиксированный итог. Лист «Объёмы»"),
        ("Пересчёт", "Если количество на объекте иное: новое количество × расценка строки"),
        ("Срок сметы", "30 календарных дней"),
        ("НДС", "Не выделен. При работе с НДС сумма считается отдельно"),
    ]
    for i, (key, value) in enumerate(meta):
        r = 10 + i
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=7)
        bg = WHITE if i % 2 == 0 else SAND
        paint(ws, r, 2, 7, fill(bg), height=20, bd=thin_border(LINE_SOFT))
        ws.cell(r, 2, key).font = font(size=10, bold=True, color=NAVY)
        ws.cell(r, 2).alignment = align("left", "center")
        ws.cell(r, 4, value).font = font(size=10, color=INK)
        ws.cell(r, 4).alignment = align("left", "center")

    ws.merge_cells("B19:G19")
    paint(ws, 19, 2, 7, fill(NAVY), height=18)
    ws["B19"] = "ИТОГО РАБОТЫ"
    ws["B19"].font = font(size=10, bold=True, color=GOLD)
    ws["B19"].alignment = align("center")
    ws["B19"].fill = fill(NAVY)

    ws.merge_cells("B20:G20")
    paint(ws, 20, 2, 7, fill(NAVY), height=32)
    ws["B20"] = f"=Смета!{totals['total_cell']}"
    ws["B20"].font = font(size=26, bold=True, color=WHITE)
    ws["B20"].alignment = align("center")
    ws["B20"].number_format = '#,##0" ₽"'
    ws["B20"].fill = fill(NAVY)

    ws.merge_cells("B21:G21")
    paint(ws, 21, 2, 7, fill(NAVY), height=18)
    ws["B21"] = "без материалов и оборудования"
    ws["B21"].font = font(size=9, color=GOLD_SOFT)
    ws["B21"].alignment = align("center")
    ws["B21"].fill = fill(NAVY)

    ws.merge_cells("B22:G22")
    ws["B22"] = rub_words(TARGET)
    ws["B22"].font = font(size=11, italic=True, color=STEEL)
    ws["B22"].alignment = align("center")
    ws["B22"].fill = fill(CREAM)
    ws.row_dimensions[22].height = 22

    ws.merge_cells("B24:G24")
    ws["B24"] = "Разделы"
    ws["B24"].font = font(size=12, bold=True, color=NAVY)
    ws["B24"].fill = fill(CREAM)
    ws.row_dimensions[24].height = 20

    for i, section in enumerate(totals["sections"]):
        r = 25 + i
        bg = WHITE if i % 2 == 0 else PALE_BLUE
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
        paint(ws, r, 2, 7, fill(bg), height=22, bd=thin_border(LINE_SOFT))
        ws.cell(r, 2, f"{section['num']}.  {section['title']}")
        ws.cell(r, 2).font = font(size=10, color=INK)
        ws.cell(r, 2).alignment = align("left", "center")
        ws.cell(r, 6, f"=Смета!{section['cell']}")
        ws.cell(r, 6).font = font(size=12, bold=True, color=NAVY)
        ws.cell(r, 6).number_format = '#,##0" ₽"'
        ws.cell(r, 6).alignment = align("right", "center")

    included = [
        "Вмазка коробок, штробление и прокладка проводов",
        "Демонтаж 2-го щита, прозвонка и распайка коробок",
        "Грунтовка, шпаклёвка в 2 слоя и вышкуривание стен 1 этажа",
        "Откосы дверного блока и монтаж дубовых подоконников",
        "Короб над печкой",
        "Трубы до котла, коллекторы с фильтром, обвязка бойлера",
    ]
    excluded = [
        "Материалы, кабель, трубы, шпаклёвка, грунт, подоконники",
        "Заделка штраб штукатуркой",
        "Покраска, обои, масло и лак подоконников",
        "Монтаж, обвязка и пуск котла",
        "Электроподключение бойлера и сборка нового щита",
        "Потолки, второй этаж и вывоз мусора",
    ]
    ws.merge_cells("B30:D30")
    ws.merge_cells("E30:G30")
    paint(ws, 30, 2, 4, fill(GREEN), height=18)
    paint(ws, 30, 5, 7, fill(TERRACOTTA), height=18)
    ws["B30"] = "ВХОДИТ"
    ws["B30"].font = font(size=10, bold=True, color=WHITE)
    ws["B30"].alignment = align("center")
    ws["B30"].fill = fill(GREEN)
    ws["E30"] = "НЕ ВХОДИТ"
    ws["E30"].font = font(size=10, bold=True, color=WHITE)
    ws["E30"].alignment = align("center")
    ws["E30"].fill = fill(TERRACOTTA)

    for i in range(6):
        r = 31 + i
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=7)
        left_bg = WHITE if i % 2 == 0 else "F3F7F4"
        right_bg = WHITE if i % 2 == 0 else "F8F1EC"
        paint(ws, r, 2, 4, fill(left_bg), height=28)
        paint(ws, r, 5, 7, fill(right_bg), height=28)
        ws.cell(r, 2, f"{i + 1}.  {included[i]}").font = font(size=8, color=INK)
        ws.cell(r, 2).alignment = align("left", "center")
        ws.cell(r, 5, f"{i + 1}.  {excluded[i]}").font = font(size=8, color=INK)
        ws.cell(r, 5).alignment = align("left", "center")

    ws.merge_cells("B38:C38")
    ws["B38"] = "Заказчик"
    ws["B38"].font = font(size=10, bold=True, color=NAVY)
    ws["B38"].fill = fill(CREAM)
    ws.merge_cells("E38:G38")
    ws["E38"] = "Исполнитель"
    ws["E38"].font = font(size=10, bold=True, color=NAVY)
    ws["E38"].fill = fill(CREAM)
    ws.row_dimensions[38].height = 18

    ws.merge_cells("B39:C39")
    ws["B39"] = "______________________________"
    ws["B39"].font = font(size=10, color=NAVY)
    ws["B39"].fill = fill(CREAM)
    ws.merge_cells("E39:G39")
    ws["E39"] = "______________________________"
    ws["E39"].font = font(size=10, color=NAVY)
    ws["E39"].fill = fill(CREAM)

    ws.merge_cells("B40:C40")
    ws["B40"] = "подпись, расшифровка"
    ws["B40"].font = font(size=8, italic=True, color=MUTED)
    ws["B40"].fill = fill(CREAM)
    ws.merge_cells("E40:G40")
    ws["E40"] = "подпись, расшифровка"
    ws["E40"].font = font(size=8, italic=True, color=MUTED)
    ws["E40"].fill = fill(CREAM)

    ws.merge_cells("B42:G42")
    ws["B42"] = (
        "Смета составлена по перечню работ на объект Симагино и фиксирует стоимость работ 230 000 ₽. "
        "Замер помещений в документ не входил: объёмы на листе «Объёмы» договорные."
    )
    ws["B42"].font = font(size=8, italic=True, color=MUTED)
    ws["B42"].alignment = align("left", "center")
    ws["B42"].fill = fill(CREAM)
    ws.row_dimensions[42].height = 28

    ws.print_area = "A1:H43"
    ws.page_setup.fitToHeight = 1
    ws.page_setup.fitToWidth = 1


def build_volumes(wb: Workbook) -> None:
    ws = wb.create_sheet("Объёмы")
    apply_print(ws, landscape=True, title="Договорные объёмы")
    set_widths(ws, {"A": 6, "B": 62, "C": 18, "D": 16, "E": 52})
    cream_sheet(ws, 24, 5)

    paint(ws, 1, 1, 5, fill(NAVY_DEEP), height=24)
    ws.merge_cells("A1:E1")
    ws["A1"] = "ДОГОВОРНЫЕ ОБЪЁМЫ   ·   ПОД ИТОГ 230 000 ₽"
    ws["A1"].font = font(size=14, bold=True, color=GOLD)
    ws["A1"].alignment = align("left", "center")
    ws["A1"].fill = fill(NAVY_DEEP)
    paint(ws, 2, 1, 5, fill(GOLD), height=5)

    ws.merge_cells("A3:E3")
    ws["A3"] = (
        "Количество ниже — основание цены, а не результат обмера. "
        "На объекте Симагино строка пересчитывается, если фактический объём отличается: новое количество × расценка листа «Смета»."
    )
    ws["A3"].font = font(size=10, color=STEEL)
    ws["A3"].alignment = align("left", "center")
    ws["A3"].fill = fill(CREAM)
    ws.row_dimensions[3].height = 32

    headers = ["№", "Показатель", "Объём", "Пункт сметы", "Пояснение"]
    for col, text in enumerate(headers, 1):
        cell = ws.cell(5, col, text)
        cell.font = font(size=10, bold=True, color=WHITE)
        cell.fill = fill(NAVY)
        cell.alignment = align("center")
        cell.border = thin_border(NAVY)
    ws.row_dimensions[5].height = 22

    for i, (name, qty, point, note) in enumerate(VOLUMES):
        r = 6 + i
        bg = WHITE if i % 2 == 0 else ROW_ALT
        ws.row_dimensions[r].height = 22
        values = [i + 1, name, qty, point, note]
        aligns = ["center", "left", "center", "center", "left"]
        for col, (value, how) in enumerate(zip(values, aligns), 1):
            cell = ws.cell(r, col, value)
            cell.font = font(size=10, bold=(col == 3), color=INK)
            cell.fill = fill(bg)
            cell.alignment = align(how)
            cell.border = thin_border(LINE_SOFT)

    ws.print_area = "A1:E17"
    ws.page_setup.fitToHeight = 1


def build_conditions(wb: Workbook) -> None:
    ws = wb.create_sheet("Условия")
    apply_print(ws, landscape=False, title="Условия")
    set_widths(ws, {"A": 4, "B": 8, "C": 78, "D": 4})
    cream_sheet(ws, 36, 4)

    paint(ws, 1, 1, 4, fill(NAVY_DEEP), height=24)
    ws.merge_cells("B1:C1")
    ws["B1"] = "УСЛОВИЯ СМЕТЫ"
    ws["B1"].font = font(size=14, bold=True, color=GOLD)
    ws["B1"].alignment = align("left", "center")
    ws["B1"].fill = fill(NAVY_DEEP)
    paint(ws, 2, 1, 4, fill(GOLD), height=5)

    blocks = [
        (
            "Цена",
            [
                "Итог 230 000 ₽ — стоимость работ по 14 пунктам перечня.",
                "Материалы, оборудование, доставка и подъём в эту сумму не входят.",
                "НДС не выделен.",
            ],
        ),
        (
            "Объёмы",
            [
                "Стены первого этажа приняты 120 м² — одна и та же площадь на грунт, два слоя шпаклёвки и шлифовку.",
                "Штраба с проводом — 60 м.п. Вмазка — 40 коробок. Распайка — 20 коробок. Прозвонка — 40 линий.",
                "Откосы — 6 м.п. одного дверного блока. Дубовые подоконники — 4 шт. Короб над печкой — 1.",
                "Трасса до котла — 15 м.п. Коллекторы с фильтром — 1 узел. Обвязка бойлера — 1.",
            ],
        ),
        (
            "Границы работ",
            [
                "Штробление включает резку и укладку провода. Заделка штрабы штукатуркой не входит.",
                "Демонтаж касается второго щита. Сборка и установка нового щита не входят.",
                "Шпаклёвка — два слоя под обои или окраску. Сама окраска и обои не входят.",
                "Подоконники монтируются без покрытия маслом или лаком.",
                "Короб над печкой — каркас и обшивка, без чистовой отделки.",
                "Трубы доводятся до котла. Сам котёл не ставится и не запускается.",
                "Обвязка бойлера — по воде. Электрическое подключение бойлера не входит.",
                "Потолки, второй этаж и вывоз мусора не входят.",
            ],
        ),
        (
            "Порядок",
            [
                "Скрытые линии прозваниваются и распаиваются до закрытия коробок.",
                "Срок действия расценок — 30 дней с даты сметы.",
                "Оплата и календарный график фиксируются в договоре, не в этой смете.",
            ],
        ),
    ]
    row = 4
    number = 1
    for title, lines in blocks:
        paint(ws, row, 2, 3, fill(STEEL), height=20)
        ws.cell(row, 2, title).font = font(size=11, bold=True, color=WHITE)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        ws.cell(row, 2).alignment = align("left", "center")
        ws.cell(row, 2).fill = fill(STEEL)
        ws.cell(row, 3).fill = fill(STEEL)
        row += 1
        for line in lines:
            bg = WHITE if number % 2 else SAND
            paint(ws, row, 2, 3, fill(bg), height=26, bd=thin_border(LINE_SOFT))
            ws.cell(row, 2, f"{number}.").font = font(size=10, bold=True, color=NAVY)
            ws.cell(row, 2).alignment = align("center")
            ws.cell(row, 3, line).font = font(size=10, color=INK)
            ws.cell(row, 3).alignment = align("left", "center")
            number += 1
            row += 1
        row += 1

    ws.print_area = f"A1:D{row}"
    ws.page_setup.fitToHeight = 1


def build_materials(wb: Workbook) -> None:
    ws = wb.create_sheet("Материалы заказчика")
    apply_print(ws, landscape=False, title="Материалы заказчика")
    set_widths(ws, {"A": 6, "B": 22, "C": 78})
    cream_sheet(ws, 22, 3)

    paint(ws, 1, 1, 3, fill(NAVY_DEEP), height=24)
    ws.merge_cells("A1:C1")
    ws["A1"] = "МАТЕРИАЛЫ ЗАКАЗЧИКА   ·   В 230 000 ₽ НЕ ВХОДЯТ"
    ws["A1"].font = font(size=13, bold=True, color=GOLD)
    ws["A1"].alignment = align("left", "center")
    ws["A1"].fill = fill(NAVY_DEEP)
    paint(ws, 2, 1, 3, fill(GOLD), height=5)

    ws.merge_cells("A3:C3")
    ws["A3"] = "Список для закупки. Цены материалов здесь нет: смета считает только работу."
    ws["A3"].font = font(size=10, italic=True, color=MUTED)
    ws["A3"].alignment = align("left", "center")
    ws["A3"].fill = fill(CREAM)
    ws.row_dimensions[3].height = 20

    for col, text in enumerate(["№", "Группа", "Что подготовить"], 1):
        cell = ws.cell(5, col, text)
        cell.font = font(size=10, bold=True, color=WHITE)
        cell.fill = fill(NAVY)
        cell.alignment = align("center")
        cell.border = thin_border(NAVY)
    ws.row_dimensions[5].height = 22

    for i, (group, item) in enumerate(MATERIALS):
        r = 6 + i
        bg = WHITE if i % 2 == 0 else ROW_ALT
        ws.row_dimensions[r].height = 22
        for col, value in enumerate([i + 1, group, item], 1):
            cell = ws.cell(r, col, value)
            cell.font = font(size=10, bold=(col == 2), color=INK)
            cell.fill = fill(bg)
            cell.alignment = align("center" if col < 3 else "left")
            cell.border = thin_border(LINE_SOFT)

    last = 6 + len(MATERIALS)
    ws.merge_cells(start_row=last + 1, start_column=1, end_row=last + 1, end_column=3)
    ws.cell(last + 1, 1, "Инструмент исполнителя (штроборез, шлифмашина, паяльник, пресс) заказчик не покупает.")
    ws.cell(last + 1, 1).font = font(size=9, italic=True, color=MUTED)
    ws.cell(last + 1, 1).alignment = align("left", "center")
    ws.cell(last + 1, 1).fill = fill(CREAM)
    ws.row_dimensions[last + 1].height = 22
    ws.print_area = f"A1:C{last + 1}"
    ws.page_setup.fitToHeight = 1


def assert_total() -> None:
    total = grand_total()
    if total != TARGET:
        details = ", ".join(f"{s['num']}={section_sum(s)}" for s in SECTIONS)
        raise SystemExit(f"Итог {total} не равен {TARGET}. Разделы: {details}")
    for section in SECTIONS:
        for item in section["items"]:
            amount = round_rub(item["qty"] * item["price"])
            if amount != item["qty"] * item["price"]:
                raise SystemExit(f"Нецелая сумма у {item['code']}")


def verify_workbook(path: Path) -> None:
    """Перечитывает файл и проверяет, что формулы дают 230 000 ₽."""
    from openpyxl import load_workbook

    book = load_workbook(path)
    sheet = book["Смета"]
    line_total = 0
    codes = []
    for row in range(1, sheet.max_row + 1):
        code = sheet.cell(row, 1).value
        if not (isinstance(code, str) and len(code) >= 3 and code[0].isdigit() and "." in code):
            continue
        qty = sheet.cell(row, 4).value
        price = sheet.cell(row, 5).value
        formula = sheet.cell(row, 6).value
        if formula != f"=ROUND(D{row}*E{row},0)":
            raise SystemExit(f"Строка {code}: формула {formula}")
        if not isinstance(qty, (int, float)) or not isinstance(price, (int, float)):
            raise SystemExit(f"Строка {code}: количество или цена не число")
        line_total += round_rub(qty * price)
        codes.append(code)
    if codes != [item["code"] for section in SECTIONS for item in section["items"]]:
        raise SystemExit(f"Состав строк не совпал: {codes}")
    if line_total != TARGET:
        raise SystemExit(f"Сумма строк {line_total} != {TARGET}")

    expected_names = [item["name"] for section in SECTIONS for item in section["items"]]
    got_names = []
    for row in range(1, sheet.max_row + 1):
        code = sheet.cell(row, 1).value
        if isinstance(code, str) and code in codes:
            got_names.append(sheet.cell(row, 2).value)
    if got_names != expected_names:
        raise SystemExit(f"Наименования не совпали: {got_names}")

    section_cells = []
    for row in range(1, sheet.max_row + 1):
        value = sheet.cell(row, 6).value
        if isinstance(value, str) and value.startswith("=SUM("):
            section_cells.append(f"F{row}")
            start, end = value[5:-1].split(":")
            chunk = 0
            for item_row in range(int(start[1:]), int(end[1:]) + 1):
                chunk += round_rub(sheet.cell(item_row, 4).value * sheet.cell(item_row, 5).value)
            if chunk <= 0:
                raise SystemExit(f"Пустой раздел в {value}")
    if len(section_cells) != 4:
        raise SystemExit(f"Ожидались 4 итога разделов, найдено {section_cells}")

    total_formula = None
    for row in range(1, sheet.max_row + 1):
        value = sheet.cell(row, 6).value
        if isinstance(value, str) and value.startswith("=") and "SUM" not in value and "ROUND" not in value:
            total_formula = value
            break
    expected_formula = "=" + "+".join(section_cells)
    if total_formula != expected_formula:
        raise SystemExit(f"Итог {total_formula} != {expected_formula}")
    if rub_words(TARGET) != "Двести тридцать тысяч рублей 00 копеек":
        raise SystemExit(rub_words(TARGET))
    if "Титул" not in book.sheetnames or book["Титул"]["B6"].value != "Объект Симагино":
        raise SystemExit("Титул не про объект Симагино")
    print("CHECK OK", line_total, "lines", len(codes))


def main() -> None:
    assert_total()
    wb = Workbook()
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True)
    totals = build_estimate(wb)
    build_summary(wb, totals)
    build_cover(wb, totals)
    build_volumes(wb)
    build_conditions(wb)
    build_materials(wb)

    order = ["Титул", "Сводка", "Смета", "Объёмы", "Условия", "Материалы заказчика"]
    for index, name in enumerate(order):
        current = wb.sheetnames.index(name)
        wb.move_sheet(name, offset=index - current)

    colors = {
        "Титул": GOLD,
        "Сводка": STEEL,
        "Смета": NAVY,
        "Объёмы": GREEN,
        "Условия": TERRACOTTA,
        "Материалы заказчика": GOLD_SOFT,
    }
    for name, color in colors.items():
        wb[name].sheet_properties.tabColor = color

    wb.properties.title = "Смета работ · Симагино"
    wb.properties.subject = "Только работы, 230 000 руб."
    wb.properties.creator = "Локальная смета"
    wb.properties.description = (
        "Объект Симагино. Электромонтаж, стены 1 этажа, откосы, подоконники, "
        "короб над печкой, трубы до котла, коллекторы и обвязка бойлера. Только работы."
    )
    wb.active = wb["Титул"]
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)

    print("FILE", OUT)
    print("TOTAL", totals["total"])
    print("WORDS", rub_words(totals["total"]))
    for section in totals["sections"]:
        print(f"SEC {section['num']} {section['sum']} {section['cell']} {section['title']}")
    verify_workbook(OUT)


if __name__ == "__main__":
    main()
